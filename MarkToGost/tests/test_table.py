# tests/test_table.py

from docx import Document

from MarkToGost.config import DocumentSettings
from MarkToGost.parser.blocks import TableBlock
from MarkToGost.renderer.document_renderer import DocumentRenderer
from MarkToGost.renderer.table import _column_widths_cm, _table_font_size

USABLE_CM = 17.0


# ================================
# Ширина колонок и размер шрифта
# ================================

class TestColumnWidths:

    def test_widths_fill_usable_width(self):
        rows = [["Код", "Название"], ["1", "Полилактид"]]
        widths = _column_widths_cm(rows, USABLE_CM, 12)
        assert abs(sum(widths) - USABLE_CM) < 1e-6

    def test_long_text_column_is_wider(self):
        rows = [["Код", "Описание"], ["1", "очень длинное описание материала для печати"]]
        code, description = _column_widths_cm(rows, USABLE_CM, 12)
        assert description > code

    def test_short_column_fits_longest_word(self):
        rows = [["№", "Текст"], ["24", "x" * 200]]
        number_width = _column_widths_cm(rows, USABLE_CM, 12)[0]
        assert number_width >= 2 * 12 * 0.021  # два символа без переноса

    def test_font_size_unchanged_for_narrow_table(self):
        rows = [["Код", "Название"], ["1", "PLA"]]
        assert _table_font_size(rows, USABLE_CM) == DocumentSettings.TABLE_FONT_SIZE_PT

    def test_font_size_reduced_for_wide_table(self):
        rows = [["Заголовокколонки"] * 12, ["значениезначение"] * 12]
        assert _table_font_size(rows, USABLE_CM) < DocumentSettings.TABLE_FONT_SIZE_PT


# ================================
# Рендеринг
# ================================

class TestRenderTable:

    def render(self, rows, caption="Материалы"):
        doc = Document()
        renderer = DocumentRenderer(doc)
        renderer.render_block(TableBlock(rows=rows, caption=caption))
        return doc, renderer

    def test_table_is_rendered_whole(self):
        rows = ["| Код | Название |"] + [f"| {i} | строка {i} |" for i in range(100)]
        doc, _ = self.render(rows)
        assert len(doc.tables) == 1
        assert len(doc.tables[0].rows) == 101

    def test_caption_numbered_and_kept_with_table(self):
        doc, renderer = self.render(["| A | B |", "| 1 | 2 |"])
        caption = doc.paragraphs[-1]
        assert caption.text == "Таблица 1 — Материалы"
        assert caption.paragraph_format.keep_with_next is True
        assert renderer.numbering.next("таблица") == "2"

    def test_table_without_caption_not_numbered(self):
        doc, renderer = self.render(["| Заказчик | Исполнитель |", "| ООО «А» | ИП Б |"], caption=None)
        assert not any(p.text.startswith("Таблица") for p in doc.paragraphs)
        assert renderer.numbering.next("таблица") == "1"

    def test_header_row_repeats(self):
        doc, _ = self.render(["| A | B |", "| 1 | 2 |"])
        tr_pr = doc.tables[0].rows[0]._tr.trPr
        assert tr_pr is not None
        assert any(el.tag.endswith("tblHeader") for el in tr_pr)
