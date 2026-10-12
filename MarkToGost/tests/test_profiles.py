# tests/test_profiles.py

import os

import pytest
from docx import Document
from docx.oxml.ns import qn

from MarkToGost.config import DocumentSettings
from MarkToGost.main import create_document
from MarkToGost.parser.blocks import HeadingBlock, Section, TextBlock
from MarkToGost.parser.frontmatter import FrontMatterError, split_front_matter
from MarkToGost.profiles import (
    ProfileError, TITLES_DIR, available_profiles, load_profile, merge, split_settings
)
from MarkToGost.parser.blocks import SheetBlock
from MarkToGost.profiles import SheetSpec, parse_sheet
from MarkToGost.renderer.sheets import field_value, fill_template, template_fields
from MarkToGost.structure import check_fields, check_structure, group_sections, insert_toc, place_sheets


# ================================
# Шапка документа
# ================================

class TestFrontMatter:

    def test_no_front_matter(self):
        assert split_front_matter("# ВВЕДЕНИЕ") == ({}, "# ВВЕДЕНИЕ")

    def test_front_matter_parsed_and_removed(self):
        data, body = split_front_matter("---\nтип: курсовая\nтема: База данных\n---\n# ВВЕДЕНИЕ\n")
        assert data == {"тип": "курсовая", "тема": "База данных"}
        assert body == "# ВВЕДЕНИЕ\n"

    def test_invalid_front_matter(self):
        with pytest.raises(FrontMatterError):
            split_front_matter("---\n- список\n---\nтекст")


# ================================
# Профили
# ================================

class TestProfiles:

    def test_all_profiles_load(self):
        keys = {p.key for p in available_profiles()}
        assert {"gost", "kursovaya", "vkr", "praktika", "referat", "dogovor"} <= keys

    def test_load_by_russian_alias(self):
        assert load_profile("Курсовая работа").key == "kursovaya"
        assert load_profile("диплом").key == "vkr"

    def test_default_profile(self):
        assert load_profile(None).key == "gost"

    def test_unknown_type(self):
        with pytest.raises(ProfileError):
            load_profile("неизвестный тип")

    def test_inherited_formatting(self):
        profile = load_profile("курсовая")
        assert profile.formatting["шрифт"] == "Times New Roman"   # из профиля «гост»
        assert profile.title_template == "kursovaya.docx"

    def test_override_in_child_profile(self):
        profile = load_profile("договор")
        assert profile.formatting["размер_шрифта"] == 12
        assert profile.toc is False
        assert profile.sections_new_page is False

    def test_title_templates_exist(self):
        for profile in available_profiles():
            if profile.title_template:
                assert os.path.isfile(os.path.join(TITLES_DIR, profile.title_template))

    def test_current_year_default(self):
        assert load_profile("курсовая").fields["год"].isdigit()

    def test_merge_is_recursive(self):
        assert merge({"a": {"x": 1, "y": 2}}, {"a": {"y": 3}}) == {"a": {"x": 1, "y": 3}}

    def test_split_settings(self):
        formatting, fields = split_settings({
            "тип": "курсовая", "тема": "Т", "размер_шрифта_кода": 9, "оформление": {"интервал": 1.0}})
        assert formatting == {"интервал": 1.0, "размер_шрифта_кода": 9}
        assert fields == {"тема": "Т"}


# ================================
# Структура
# ================================

def heading(text):
    return HeadingBlock(text=text, level=1, marker=1)


class TestStructure:

    def test_group_sections(self):
        blocks = [heading("ВВЕДЕНИЕ"), TextBlock("а"), heading("1 АНАЛИЗ"), TextBlock("б")]
        sections = group_sections(blocks, new_page=True)
        assert [s.section_id for s in sections] == ["ВВЕДЕНИЕ", "1 АНАЛИЗ"]
        assert sections[0].blocks == [TextBlock("а")]

    def test_subheadings_stay_inside_section(self):
        blocks = [heading("1 АНАЛИЗ"), HeadingBlock(text="1.1 Обзор", level=2, marker=2)]
        assert len(group_sections(blocks, new_page=True)[0].blocks) == 1

    def test_toc_inserted_after_before_toc_items(self):
        profile = load_profile("вкр")
        blocks = group_sections([heading("РЕФЕРАТ"), heading("ВВЕДЕНИЕ")], new_page=True)
        result = insert_toc(blocks, profile)
        assert [s.section_id for s in result] == ["РЕФЕРАТ", "СОДЕРЖАНИЕ", "ВВЕДЕНИЕ"]
        assert result[1].toc is True

    def test_existing_toc_kept(self):
        profile = load_profile("курсовая")
        blocks = [Section(section_id="ОГЛАВЛЕНИЕ", blocks=[], toc=True)]
        assert insert_toc(blocks, profile) == blocks

    def test_complete_structure_has_no_warnings(self):
        profile = load_profile("курсовая")
        titles = ["ВВЕДЕНИЕ", "1 АНАЛИЗ", "ЗАКЛЮЧЕНИЕ", "Список литературы", "ПРИЛОЖЕНИЕ А"]
        blocks = group_sections([heading(t) for t in titles], new_page=True)
        assert check_structure(blocks, profile) == []

    def test_missing_section_reported(self):
        profile = load_profile("курсовая")
        blocks = group_sections([heading("ВВЕДЕНИЕ"), heading("1 АНАЛИЗ")], new_page=True)
        warnings = check_structure(blocks, profile)
        assert any("ЗАКЛЮЧЕНИЕ" in w for w in warnings)
        assert any("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ" in w for w in warnings)

    def test_wrong_order_reported(self):
        profile = load_profile("курсовая")
        titles = ["ЗАКЛЮЧЕНИЕ", "ВВЕДЕНИЕ", "1 АНАЛИЗ", "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ"]
        blocks = group_sections([heading(t) for t in titles], new_page=True)
        assert any("должен идти после" in w for w in check_structure(blocks, profile))

    def test_required_fields(self):
        warnings = check_fields({"тема": "Т"}, ["тема", "студент"], "задание")
        assert warnings == ["не заполнено поле «студент» (задание)"]


# ================================
# Титульный лист
# ================================

class TestTitlePage:

    def text(self, doc):
        return "\n".join(p.text for p in doc.paragraphs)

    def test_fields_substituted(self):
        fields = {"министерство": "М", "вуз": "Университет", "факультет": "", "кафедра": "ПО",
                  "вид_работы": "КУРСОВАЯ РАБОТА", "дисциплина": "Базы данных", "тема": "Тема",
                  "студент": "Иванов И. И.", "группа": "ПИ-1", "должность_руководителя": "",
                  "руководитель": "Петров П. П.", "город": "Москва", "год": "2026"}
        doc, unknown = fill_template(os.path.join(TITLES_DIR, "kursovaya.docx"), fields)
        text = self.text(doc)
        assert unknown == []
        assert "на тему «Тема»" in text
        assert "Иванов И. И." in text
        assert "Москва 2026" in text
        assert "{{" not in text

    def test_line_with_empty_field_removed(self):
        doc, _ = fill_template(os.path.join(TITLES_DIR, "kursovaya.docx"), {"вуз": "У"})
        text = self.text(doc)
        assert "Кафедра" not in text   # поле «кафедра» не задано — строки нет
        assert "{{" not in text


    def test_multiline_and_list_values(self):
        fields = {"разделы_записки": ["Введение", "1 Анализ", "Заключение"], "тема": "Т"}
        doc, _ = fill_template(os.path.join(TITLES_DIR, "zadanie.docx"), fields)
        p = next(p for p in doc.paragraphs if "Введение" in p.text)
        assert p.text.endswith("Введение\n1 Анализ\nЗаключение")

    def test_field_value(self):
        assert field_value(["а", "б"]) == "а\nб"
        assert field_value(None) == ""
        assert field_value("текст\n") == "текст"

    def test_template_fields(self):
        names = template_fields(os.path.join(TITLES_DIR, "zadanie.docx"))
        assert {"тема", "исходные_данные", "разделы_записки", "заведующий"} <= set(names)


# ================================
# Размещение вставных листов
# ================================

def sections(*titles):
    return [Section(section_id=t, blocks=[]) for t in titles]


class TestPlaceSheets:

    def spec(self, name, **place):
        return parse_sheet(name, {"шаблон": f"{name}.docx", **place})

    def names(self, blocks):
        return [b.name if isinstance(b, SheetBlock) else b.section_id for b in blocks]

    def test_title_then_sheet_after_title(self):
        title = SheetSpec(name="титульный лист", template="t.docx", place="начало")
        blocks, warnings = place_sheets(sections("ВВЕДЕНИЕ"), [title, self.spec("задание", после="титульный лист")])
        assert self.names(blocks) == ["титульный лист", "задание", "ВВЕДЕНИЕ"]
        assert warnings == []

    def test_before_and_after_section(self):
        blocks, _ = place_sheets(sections("ВВЕДЕНИЕ", "ЗАКЛЮЧЕНИЕ", "ПРИЛОЖЕНИЕ А"), [
            self.spec("отзыв", перед="приложение"), self.spec("справка", после="ВВЕДЕНИЕ")])
        assert self.names(blocks) == ["ВВЕДЕНИЕ", "справка", "ЗАКЛЮЧЕНИЕ", "отзыв", "ПРИЛОЖЕНИЕ А"]

    def test_end_and_order_kept(self):
        blocks, _ = place_sheets(sections("ВВЕДЕНИЕ"), [self.spec("а", место="конец"), self.spec("б", место="конец")])
        assert self.names(blocks) == ["ВВЕДЕНИЕ", "а", "б"]

    def test_after_toc(self):
        toc = Section(section_id="СОДЕРЖАНИЕ", blocks=[], toc=True)
        blocks, _ = place_sheets([toc] + sections("ВВЕДЕНИЕ"), [self.spec("перечень", после="содержание")])
        assert self.names(blocks) == ["СОДЕРЖАНИЕ", "перечень", "ВВЕДЕНИЕ"]

    def test_missing_anchor_reported(self):
        blocks, warnings = place_sheets(sections("ВВЕДЕНИЕ"), [self.spec("отзыв", перед="ЗАКЛЮЧЕНИЕ")])
        assert self.names(blocks) == ["ВВЕДЕНИЕ"]
        assert warnings == ["лист «отзыв» не вставлен: нет раздела «ЗАКЛЮЧЕНИЕ»"]

    def test_invalid_place(self):
        with pytest.raises(ProfileError):
            parse_sheet("x", {"шаблон": "x.docx", "место": "середина"})

    def test_profile_sheet_definition(self):
        spec = load_profile("курсовая").sheets["задание"]
        assert (spec.place, spec.anchor, spec.page_numbers) == ("после", "титульный лист", False)


# ================================
# Сборка документа по профилю
# ================================

KURSOVAYA = """---
тип: курсовая
вуз: Университет
кафедра: Программное обеспечение
дисциплина: Базы данных
тема: Тема
студент: Иванов И. И.
группа: ПИ-1
руководитель: Петров П. П.
город: Москва
размер_шрифта_кода: 9
---

# ВВЕДЕНИЕ

Текст.

# 1 АНАЛИЗ

```
SELECT 1;
```

# ЗАКЛЮЧЕНИЕ

Текст.

# СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ

1. Источник.
"""


class TestCreateDocument:

    def test_profile_document(self):
        warnings = []
        doc = create_document(KURSOVAYA, warnings=warnings)
        texts = [p.text for p in doc.paragraphs]
        assert warnings == []
        assert any("Иванов И. И." in t for t in texts)        # титульный лист
        assert "СОДЕРЖАНИЕ" in texts                           # содержание вставлено
        assert "ВВЕДЕНИЕ" in texts
        assert len(doc.sections) == 2                          # титульный лист — отдельный раздел
        code = next(p for p in doc.paragraphs if p.text == "SELECT 1;")
        assert code.runs[0].font.size.pt == 9                  # параметр из шапки

    def test_settings_restored_after_document(self):
        before = DocumentSettings.CODE_FONT_SIZE_PT
        create_document(KURSOVAYA)
        assert DocumentSettings.CODE_FONT_SIZE_PT == before

    def test_title_page_disabled(self):
        doc = create_document(KURSOVAYA.replace("тип: курсовая", "тип: курсовая\nтитульный_лист: нет"))
        assert not any("Иванов И. И." in p.text for p in doc.paragraphs)
        assert len(doc.sections) == 1

    def test_without_front_matter_unchanged(self):
        doc = create_document("# ВВЕДЕНИЕ\n\nТекст.")
        texts = [p.text for p in doc.paragraphs]
        assert "СОДЕРЖАНИЕ" not in texts
        assert len(doc.sections) == 1

    def test_assignment_sheet_inserted(self):
        md = KURSOVAYA.replace("тип: курсовая", "тип: курсовая\nлисты: [задание]\nзаведующий: Сидоров С. С.")
        warnings = []
        doc = create_document(md, warnings=warnings)
        texts = [p.text for p in doc.paragraphs]
        heading = "ЗАДАНИЕ\nна курсовую работу\nпо дисциплине «Базы данных»"
        assert heading in texts
        assert len(doc.sections) == 3                          # титульный лист, задание, основная часть
        assert texts.index(heading) < texts.index("СОДЕРЖАНИЕ")

    def test_inline_sheet_directive(self, tmp_path):
        template = Document()
        template.add_paragraph("Отзыв на работу «{{тема}}»")
        template.save(tmp_path / "отзыв.docx")
        md = KURSOVAYA + "\n{{лист: отзыв.docx | без номера}}\n"
        warnings = []
        doc = create_document(md, base_dir=str(tmp_path), warnings=warnings)
        assert "Отзыв на работу «Тема»" in [p.text for p in doc.paragraphs]
        assert warnings == []

    def test_unknown_profile_sheet_reported(self):
        warnings = []
        create_document(KURSOVAYA.replace("тип: курсовая", "тип: курсовая\nлисты: [отзыв]"), warnings=warnings)
        assert any("нет листа «отзыв»" in w for w in warnings)

    def test_title_page_has_no_page_number(self):
        doc = create_document(KURSOVAYA)
        title_sectPr = doc.sections[0]._sectPr
        assert title_sectPr.find(qn("w:titlePg")) is not None
