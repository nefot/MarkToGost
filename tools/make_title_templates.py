#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Создаёт стартовые шаблоны титульных листов в MarkToGost/profiles/titles/.

Шаблон — обычный DOCX с полями {{имя}}; значения подставляются из шапки документа.
Строка шаблона, в которой поле оказалось пустым, удаляется целиком. Готовые шаблоны можно
переделать в Word под требования своего вуза — повторно запускать этот скрипт не нужно.

Использование:
    python tools/make_title_templates.py           # создать недостающие шаблоны
    python tools/make_title_templates.py --force   # пересоздать все (правки в Word будут потеряны)
"""

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from MarkToGost.profiles import TITLES_DIR  # noqa: E402

FONT = "Times New Roman"
SIGN = "____________"


def new_document(v_align="both"):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin, sec.right_margin = Cm(3.0), Cm(1.0)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
    sectPr = sec._sectPr
    # Блоки титульного листа распределяются по высоте страницы
    if v_align:
        el = OxmlElement("w:vAlign")
        el.set(qn("w:val"), v_align)
        sectPr.append(el)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(14)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    # Пустой абзац шаблона по умолчанию не нужен
    body = doc.element.body
    for p in list(body.iterchildren(qn("w:p"))):
        body.remove(p)
    return doc


def block(doc, lines, align=WD_ALIGN_PARAGRAPH.CENTER, left_indent_cm=0.0, space_after=0,
          size=14, line_spacing=1.15):
    """Абзац из строк, разделённых переносом. lines — строки или [(текст, жирный, размер)]"""
    p = doc.add_paragraph()
    p.alignment = align
    fmt = p.paragraph_format
    fmt.first_line_indent = Cm(0)
    fmt.left_indent = Cm(left_indent_cm)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = line_spacing
    default_size = size
    for i, line in enumerate(lines):
        text, bold, size = (line, False, default_size) if isinstance(line, str) else line
        run = p.add_run(text)
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        if i < len(lines) - 1:
            run.add_break(WD_BREAK.LINE)
    return p


def header(doc):
    block(doc, [
        ("{{министерство}}", False, 12),
        ("{{вуз}}", True, 14),
        "{{факультет}}",
        "Кафедра «{{кафедра}}»",
    ])


def footer(doc):
    block(doc, ["{{город}} {{год}}"])


def signatures(doc, lines):
    block(doc, lines, align=WD_ALIGN_PARAGRAPH.LEFT, left_indent_cm=8.0)


def kursovaya(path, work="{{вид_работы}}"):
    doc = new_document()
    header(doc)
    block(doc, [
        (work, True, 16),
        "по дисциплине «{{дисциплина}}»",
        ("на тему «{{тема}}»", True, 14),
    ])
    signatures(doc, [
        "Выполнил студент группы {{группа}}",
        f"{{{{студент}}}} {SIGN}",
        "",
        "Руководитель",
        "{{должность_руководителя}}",
        f"{{{{руководитель}}}} {SIGN}",
        "",
        f"Оценка {SIGN}",
    ])
    footer(doc)
    doc.save(path)


def vkr(path):
    doc = new_document()
    header(doc)
    block(doc, [
        "ДОПУСТИТЬ К ЗАЩИТЕ",
        "Заведующий кафедрой",
        f"{{{{заведующий}}}} {SIGN}",
        "«___» ____________ {{год}} г.",
    ], align=WD_ALIGN_PARAGRAPH.LEFT, left_indent_cm=8.0)
    block(doc, [
        ("{{вид_работы}}", True, 16),
        ("на тему «{{тема}}»", True, 14),
        "Направление подготовки {{направление}}",
        "Профиль «{{профиль}}»",
    ])
    signatures(doc, [
        "Выполнил студент группы {{группа}}",
        f"{{{{студент}}}} {SIGN}",
        "",
        "Руководитель",
        "{{должность_руководителя}}",
        f"{{{{руководитель}}}} {SIGN}",
        f"Консультант {{{{консультант}}}} {SIGN}",
        f"Нормоконтроль {{{{нормоконтроль}}}} {SIGN}",
    ])
    footer(doc)
    doc.save(path)


def praktika(path):
    doc = new_document()
    header(doc)
    block(doc, [
        ("ОТЧЁТ", True, 16),
        ("о прохождении {{вид_практики}} практики", True, 14),
        "на базе {{организация}}",
        "Сроки практики: {{сроки}}",
    ])
    signatures(doc, [
        "Выполнил студент группы {{группа}}",
        f"{{{{студент}}}} {SIGN}",
        "",
        f"Руководитель от организации {{{{руководитель_от_организации}}}} {SIGN}",
        "Руководитель от кафедры",
        f"{{{{руководитель}}}} {SIGN}",
        "",
        f"Оценка {SIGN}",
    ])
    footer(doc)
    doc.save(path)


def zadanie(path):
    """Задание на курсовую работу или ВКР — по форме кафедры (лист задания), 12 pt на одной странице"""
    doc = new_document(v_align=None)
    left = WD_ALIGN_PARAGRAPH.LEFT
    small = dict(size=12, line_spacing=1.0)
    block(doc, [("{{министерство}}", False, 11), ("{{вуз}}", True, 12)], space_after=8, **small)
    block(doc, [
        "Факультет {{факультет}}",
        "Кафедра «{{кафедра}}»",
        "Направление подготовки {{направление}}",
    ], align=left, space_after=8, **small)
    block(doc, [
        "УТВЕРЖДАЮ",
        "Заведующий кафедрой",
        f"{{{{заведующий}}}} {SIGN}",
        "«___» ____________ {{год}} г.",
    ], align=left, left_indent_cm=9.0, space_after=8, **small)
    block(doc, [("ЗАДАНИЕ", True, 14), "на {{задание_на}}", "по дисциплине «{{дисциплина}}»"],
          space_after=8, **small)
    block(doc, ["Студенту {{студент}}, группа {{группа}}"], align=left, space_after=6, **small)
    for title, field in [("1. Тема работы:", "тема"),
                         ("2. Исходные данные:", "исходные_данные"),
                         ("3. Разделы пояснительной записки:", "разделы_записки"),
                         ("4. Перечень обязательного графического материала:", "графический_материал")]:
        block(doc, [(title, True, 12), "{{" + field + "}}"], align=left, space_after=6, **small)
    block(doc, [
        "Дата выдачи задания: {{дата_выдачи}}",
        "Срок сдачи законченной работы: {{срок_сдачи}}",
    ], align=left, space_after=10, **small)
    block(doc, [
        f"Руководитель {{{{руководитель}}}} {SIGN}",
        f"Задание принял к исполнению {{{{студент}}}} {SIGN}",
    ], align=left, **small)
    doc.save(path)


TEMPLATES = {
    "kursovaya.docx": kursovaya,
    "referat.docx": kursovaya,
    "vkr.docx": vkr,
    "praktika.docx": praktika,
    "zadanie.docx": zadanie,
}


def main():
    for stream in (sys.stdout, sys.stderr):  # консоль Windows (cp1251)
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")
    force = "--force" in sys.argv
    out = Path(TITLES_DIR)
    out.mkdir(exist_ok=True)
    for name, make in TEMPLATES.items():
        path = out / name
        if path.exists() and not force:
            print(f"пропущен (уже есть): {path}")
            continue
        make(path)
        print(f"создан: {path}")


if __name__ == "__main__":
    main()
